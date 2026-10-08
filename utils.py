import os
import pandas as pd
import openpyxl
import io

def find_excel_file():
    for file in os.listdir('.'):
        if file.endswith('.xlsx') and not file.startswith('~$') and not file.startswith('BaoCao_'):
            return file
    return "3. PHÂN LOẠI HỒ SƠ ĐỦ ĐIỀU KIỆN XÉT.xlsx"

def load_and_process_excel(file_path=None):
    if not file_path or not os.path.exists(file_path):
        file_path = find_excel_file()
        
    if not os.path.exists(file_path):
        return pd.DataFrame(), pd.DataFrame()

    try:
        wb = openpyxl.load_workbook(file_path, data_only=True)
    except Exception as e:
        print(f"Lỗi mở file Excel: {e}")
        return pd.DataFrame(), pd.DataFrame()

    # -------------------------------------------------------------------------
    # 1. ĐỌC SHEET "PHÂN LOẠI 74"
    # -------------------------------------------------------------------------
    df_74 = pd.DataFrame()
    if 'PHÂN LOẠI 74' in wb.sheetnames:
        ws_74 = wb['PHÂN LOẠI 74']
        data_74 = []
        current_linh_vuc = "Chưa phân loại"
        
        for r in range(1, ws_74.max_row + 1):
            stt = ws_74.cell(r, 1).value
            ten_nv = ws_74.cell(r, 2).value
            
            if stt and isinstance(str(stt), str) and "Lĩnh vực" in str(stt):
                current_linh_vuc = str(stt).strip()
                continue
                
            if isinstance(stt, (int, float)) or (isinstance(stt, str) and str(stt).strip().isdigit()):
                don_vi = ws_74.cell(r, 3).value or "Chưa xác định"
                kp_val = ws_74.cell(r, 4).value
                
                kp_so = 0.0
                if kp_val:
                    try:
                        str_kp = str(kp_val).split('\n')[0].replace(',', '.').strip()
                        kp_so = float(str_kp)
                    except:
                        kp_so = 0.0

                data_74.append({
                    "STT": int(stt),
                    "Lĩnh vực": current_linh_vuc,
                    "Tên đề xuất/nhiệm vụ": ten_nv or "",
                    "Đơn vị đề xuất": don_vi,
                    "Kinh phí đề xuất (trđ)": kp_so,
                    "Kinh phí chi tiết": str(kp_val) if kp_val else "0",
                    "Nội dung": ws_74.cell(r, 5).value or "",
                    "Loại hình nhiệm vụ": str(ws_74.cell(r, 6).value).strip() if ws_74.cell(r, 6).value else "Khác",
                    "Ghi chú": ws_74.cell(r, 7).value or ""
                })
        df_74 = pd.DataFrame(data_74)

    # -------------------------------------------------------------------------
    # 2. ĐỌC SHEET "TỔNG HỢP 25" & GỘP Ý KIẾN TỪ CỘT D, F, H, J, M SHEET "Ý KIẾN CỦA PHÒNG"
    # -------------------------------------------------------------------------
    df_25 = pd.DataFrame()
    if 'TỔNG HỢP 25' in wb.sheetnames:
        ws_25 = wb['TỔNG HỢP 25']
        ws_yk = wb['Ý KIẾN CỦA PHÒNG'] if 'Ý KIẾN CỦA PHÒNG' in wb.sheetnames else None
        
        col_yk_map = [
            (4, "Đề xuất Tên nhiệm vụ (Cột D)"),
            (6, "Đề xuất Mục tiêu (Cột F)"),
            (8, "Đề xuất Nội dung (Cột H)"),
            (10, "Đề xuất Sản phẩm (Cột J)"),
            (13, "Một số ý kiến của Phòng QLKH (Cột M)")
        ]

        data_25 = []
        for r in range(2, ws_25.max_row + 1):
            stt = ws_25.cell(r, 1).value
            if not stt or not (isinstance(stt, (int, float)) or str(stt).strip().isdigit()):
                continue
                
            don_vi = ws_25.cell(r, 2).value or ""
            ten_nv = ws_25.cell(r, 3).value or ""
            muc_tieu = ws_25.cell(r, 4).value or ""
            noi_dung = ws_25.cell(r, 5).value or ""
            san_pham = ws_25.cell(r, 6).value or ""
            kinh_phi = ws_25.cell(r, 7).value or 0
            loai_hinh = str(ws_25.cell(r, 8).value).strip() if ws_25.cell(r, 8).value else "Khác"
            ghi_chu = ws_25.cell(r, 9).value or ""
            
            ykien_gop_list = []
            if ws_yk and r <= ws_yk.max_row:
                for col_idx, label in col_yk_map:
                    val = ws_yk.cell(r, col_idx).value
                    if val is not None and str(val).strip() != "":
                        ykien_gop_list.append(f"🔹 **{label}:**\n{str(val).strip()}")

            y_kien_tong_hop = "\n\n---\n\n".join(ykien_gop_list) if ykien_gop_list else ""

            data_25.append({
                "STT": int(stt),
                "ĐƠN VỊ ĐỀ XUẤT": don_vi,
                "TÊN NHIỆM VỤ": ten_nv,
                "MỤC TIÊU": muc_tieu,
                "NỘI DUNG": noi_dung,
                "SẢN PHẨM": san_pham,
                "KINH PHÍ THỰC HIỆN (TRĐ)": kinh_phi,
                "LOẠI HÌNH NHIỆM VỤ": loai_hinh,
                "GHI CHÚ": ghi_chu,
                "MỘT SỐ Ý KIẾN CỦA PHÒNG QLKH": y_kien_tong_hop
            })
        df_25 = pd.DataFrame(data_25)

    return df_74, df_25

def generate_exact_template_excel(file_path=None):
    """Xuất file Excel đúng chuẩn nguyên bản theo file gốc, đã cập nhật cột ý kiến gộp"""
    if not file_path or not os.path.exists(file_path):
        file_path = find_excel_file()
        
    wb = openpyxl.load_workbook(file_path)
    
    if 'TỔNG HỢP 25' in wb.sheetnames and 'Ý KIẾN CỦA PHÒNG' in wb.sheetnames:
        ws_25 = wb['TỔNG HỢP 25']
        ws_yk = wb['Ý KIẾN CỦA PHÒNG']
        
        col_yk_map = [(4, "Đề xuất Tên"), (6, "Đề xuất Mục tiêu"), (8, "Đề xuất Nội dung"), (10, "Đề xuất Sản phẩm"), (13, "Ý kiến Phòng QLKH")]
        
        for r in range(2, ws_25.max_row + 1):
            stt = ws_25.cell(r, 1).value
            if not stt or not (isinstance(stt, (int, float)) or str(stt).strip().isdigit()):
                continue
                
            ykien_gop_list = []
            if r <= ws_yk.max_row:
                for col_idx, label in col_yk_map:
                    val = ws_yk.cell(r, col_idx).value
                    if val is not None and str(val).strip() != "":
                        ykien_gop_list.append(f"[{label}]: {str(val).strip()}")
            
            y_kien_text = "\n".join(ykien_gop_list) if ykien_gop_list else ""
            ws_25.cell(r, 10).value = y_kien_text  # Cột J của sheet TỔNG HỢP 25

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()